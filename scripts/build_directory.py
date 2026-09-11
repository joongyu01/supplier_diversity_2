"""Build an evidence-based supplier directory from local official workbooks.

No generated products, prices, company-name joins, or certificate inference.
"""
import argparse
import datetime as dt
import hashlib
import json
from pathlib import Path
import re
from urllib.parse import urlparse

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOCIAL_URL = 'https://www.sepp.or.kr/board/value/NOTICE/boardDtl?lettNo=278'


def clean(v):
    return re.sub(r'\s+', ' ', str(v or '')).strip()


def bizno(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r'[-\s]', '', str(v or ''))
    return s if re.fullmatch(r'\d{10}', s) else ''


def date(v):
    if isinstance(v, (dt.date, dt.datetime)):
        return v.strftime('%Y-%m-%d')
    s = str(v or '').strip()
    if re.fullmatch(r'\d{8}', s):
        s = f'{s[:4]}-{s[4:6]}-{s[6:]}'
    try:
        return dt.date.fromisoformat(s[:10]).isoformat()
    except ValueError:
        return None


def in_period(start, end, cancel, as_of):
    return not ((start and start > as_of) or (end and end < as_of)
                or (cancel and cancel <= as_of))


def website(v):
    s = clean(v)
    if not s or s in ('없음', '-', '해당없음'):
        return ''
    if '://' not in s:
        s = 'https://' + s
    p = urlparse(s)
    if p.scheme not in ('https', 'http') or not p.hostname or '.' not in p.hostname or p.username:
        return ''
    return s


def region(v):
    first = clean(v).split(' ')[0]
    aliases = {'서울특별시': '서울', '경기도': '경기', '인천광역시': '인천', '부산광역시': '부산',
               '대구광역시': '대구', '광주광역시': '광주', '대전광역시': '대전', '울산광역시': '울산',
               '세종특별자치시': '세종', '강원특별자치도': '강원', '강원도': '강원', '충청북도': '충북',
               '충청남도': '충남', '전북특별자치도': '전북', '전라북도': '전북', '전라남도': '전남',
               '경상북도': '경북', '경상남도': '경남', '제주특별자치도': '제주', '제주도': '제주'}
    return aliases.get(first, first if first in aliases.values() else '')


def evidence(kind, source, start=None, end=None):
    return {'type': kind, 'source': source, 'validFrom': start,
            'validUntil': None if end == '9999-12-31' else end}


def add_type(company, item):
    if item not in company['evidence']:
        company['evidence'].append(item)


def company(code, name):
    return {'bizno': code, 'name': clean(name), 'region': '', 'website': '',
            'business': '', 'items': [], 'industries': [], 'evidence': [],
            'certificationNo': '', 'branch': '', 'excludedAsLargeCorp': False, 'contacts': []}


def build(social_path, policy_path, as_of):
    companies = {}
    wb = openpyxl.load_workbook(social_path, read_only=True, data_only=True)
    rows = wb.active.iter_rows(values_only=True)
    next(rows)
    headers = [clean(x) for x in next(rows)]
    required = ['인증번호', '사업자번호 (본점,지점)', '인증상태', '기업명', '사업내용', '홈페이지']
    if not all(x in headers for x in required):
        raise ValueError('Official social workbook headers changed')
    social_rows = 0
    for row in rows:
        r = dict(zip(headers, row))
        code = bizno(r.get('사업자번호 (본점,지점)'))
        if not code or clean(r.get('인증상태')) != '인증':
            continue
        start, end, cancel = (date(r.get(k)) for k in ('실적시작일(인증일자)', '실적종료일(실적인정)', '인증취소일(실적불인정)'))
        if not in_period(start, end, cancel, as_of):
            continue
        social_rows += 1
        c = companies.setdefault(code, company(code, r['기업명']))
        c.update(region=region(r.get('지역명')), website=website(r.get('홈페이지')),
                 business=clean(r.get('사업내용')), certificationNo=clean(r.get('인증번호')),
                 branch=clean(r.get('본지점')), industries=[clean(r.get(k)) for k in ('대분류명', '중분류명') if r.get(k)])
        c['contacts'].append({'source': 'social', 'representative': clean(r.get('대표자')),
                              'phone': clean(r.get('대표전화번호')), 'address': clean(r.get('소재지'))})
        add_type(c, evidence('사회적기업', 'social', start, end))
    wb.close()

    wb = openpyxl.load_workbook(policy_path, read_only=True, data_only=True)
    severe = wb["중증('26.06.30)"]
    severe_codes = set()
    for r in severe.iter_rows(min_row=2, values_only=True):
        code = bizno(r[0])
        start, end = date(r[8]), date(r[9])
        if not code or not in_period(start, end, None, as_of):
            continue
        severe_codes.add(code)
        c = companies.setdefault(code, company(code, r[1]))
        c['region'] = c['region'] or region(r[4])
        c['items'] = sorted(set(c['items']) | {clean(p) for p in clean(r[7]).split(',') if clean(p)})
        contact = {'source': 'policy', 'facilityHead': clean(r[2]), 'phone': clean(r[6]), 'address': clean(r[5])}
        if contact not in c['contacts']:
            c['contacts'].append(contact)
        add_type(c, evidence('중증장애인생산품 생산시설', 'policy', start, end))

    registry_counts = {'사회적기업': social_rows, '중증장애인생산품 생산시설': len(severe_codes)}
    # Only enrich suppliers with real business/item descriptions. No fabricated offerings.
    for sheet, kind, has_cancel in [("여성기업('26.06.30)", '여성기업', True),
                                     ("장애인기업('26.06.30)", '장애인기업', True),
                                     ("창업('26.06.30)", '창업기업', False)]:
        latest = {}
        representatives = {}
        for r in wb[sheet].iter_rows(min_row=2, values_only=True):
            code = bizno(r[0])
            if not code:
                continue
            start, end = date(r[5]), date(r[6])
            cancel = date(r[7]) if has_cancel else None
            rec = (start or '', end or '', cancel or '')
            if code not in latest or rec > latest[code]:
                latest[code] = rec
                representatives[code] = clean(r[2])
        active = {code: rec for code, rec in latest.items() if in_period(*rec, as_of)}
        registry_counts[kind] = len(active)
        for code in companies.keys() & active.keys():
            start, end, _ = active[code]
            add_type(companies[code], evidence(kind, 'policy', start or None, end or None))
            rep = representatives.get(code)
            if rep and not any(c.get('representative') == rep for c in companies[code]['contacts']):
                companies[code]['contacts'].append({'source': 'policy', 'type': kind, 'representative': rep})
    for sheet, kind in [("장애인표준사업장('26.06.30", '장애인표준사업장'),
                        ("협동조합('26.06.30)", '사회적협동조합')]:
        codes = {bizno(r[0]) for r in wb[sheet].iter_rows(min_row=2, values_only=True)} - {''}
        registry_counts[kind] = len(codes)
        for code in companies.keys() & codes:
            add_type(companies[code], evidence(kind, 'policy'))
    for r in wb['대기업(26.01.12.)+상호출자'].iter_rows(min_row=2, values_only=True):
        if bizno(r[1]) in companies:
            companies[bizno(r[1])]['excludedAsLargeCorp'] = True
    wb.close()

    data = []
    for c in companies.values():
        c['types'] = sorted({e['type'] for e in c['evidence']})
        if c['business'] or c['items'] or c['industries']:
            data.append(c)
    data.sort(key=lambda c: (c['name'], c['bizno']))
    coverage = {kind: {'registryCount': count, 'searchableCount': sum(kind in c['types'] for c in data)}
                for kind, count in registry_counts.items()}
    return {'schemaVersion': 1, 'generatedAt': dt.datetime.now(dt.timezone.utc).isoformat(), 'checkDate': as_of,
            'sources': {'social': {'label': '가치장터 인증 사회적기업 명단', 'basisDate': '2026-07-14', 'url': SOCIAL_URL,
                                    'sha256': hashlib.sha256(social_path.read_bytes()).hexdigest()},
                        'policy': {'label': '정부권장정책 기업목록', 'basisDate': '2026-06-30',
                                   'file': policy_path.name, 'sha256': hashlib.sha256(policy_path.read_bytes()).hexdigest()}},
            'coverage': coverage, 'suppliers': data}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--social', type=Path, default=ROOT / 'private/sources/social-enterprises-2026-2.xlsx')
    p.add_argument('--policy', type=Path, default=ROOT / '2026년 6월 30일 기준 정부권장정책.xlsx')
    p.add_argument('--as-of', default=dt.date.today().isoformat())
    args = p.parse_args()
    dt.date.fromisoformat(args.as_of)
    result = build(args.social, args.policy, args.as_of)
    path = ROOT / 'site/data/directory.json'
    tmp = path.with_suffix('.tmp')
    tmp.write_text(json.dumps(result, ensure_ascii=False, separators=(',', ':')) + '\n', encoding='utf-8')
    tmp.replace(path)
    print(json.dumps({'suppliers': len(result['suppliers']), 'bytes': path.stat().st_size,
                      'coverage': result['coverage']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
