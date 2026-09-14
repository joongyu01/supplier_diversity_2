"""Extract complete business registries, without item or industry filters."""
import datetime as dt
import hashlib
import json
from collections import Counter
from pathlib import Path
import openpyxl
from build_directory import bizno, clean, date

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'data/registries'
ASOF=dt.date.today().isoformat()
TYPES={'sme':'중소기업','women':'여성기업','disabled':'장애인기업','startup':'창업기업','severe':'중증장애인생산품 생산시설','standard':'장애인표준사업장','social':'사회적기업','cooperative':'사회적협동조합'}

def status(start,end,cancel,explicit=''):
    if explicit and explicit!='인증': return 'source_cancelled'
    if cancel and cancel <= ASOF: return 'cancelled'
    if start and start > ASOF: return 'not_started'
    if end and end < ASOF: return 'expired'
    return 'within_recorded_period' if start or end else 'listed_date_unknown'

def run():
    OUT.mkdir(parents=True,exist_ok=True)
    previous=json.loads((OUT/'manifest.json').read_text(encoding='utf-8')) if (OUT/'manifest.json').exists() else {}
    previous_smpp=previous.get('sources',{}).get('smpp')
    if previous_smpp and not (ROOT/'private/sources/smpp'/previous_smpp['basisDate']).is_dir():
        raise FileNotFoundError('Preserved SMPP source archive is missing; keep existing registries unchanged.')
    registries={k:{} for k in TYPES if k!='sme'}
    invalid=Counter(); rows=Counter(); master={}
    def add(kind,code,name,source,sheet,row,representative='',phone='',address='',start=None,end=None,cancel=None,explicit='',**extra):
        rows[kind]+=1
        code=bizno(code)
        if not code: invalid[kind]+=1; return
        rec={'source':source,'sheet':sheet,'row':row,'name':clean(name),'representative':clean(representative),'phone':clean(phone),'address':clean(address),'validFrom':date(start),'validUntil':date(end),'cancelledAt':date(cancel),**extra}
        rec['status']=status(rec['validFrom'],rec['validUntil'],rec['cancelledAt'],explicit)
        c=registries[kind].setdefault(code,{'bizno':code,'type':TYPES[kind],'records':[]})
        c['records'].append(rec)
    policy=ROOT/'2026년 6월 30일 기준 정부권장정책.xlsx'
    social=ROOT/'private/sources/social-enterprises-2026-2.xlsx'
    w=openpyxl.load_workbook(policy,read_only=True,data_only=True)
    mapping=[('women',"여성기업('26.06.30)"),('disabled',"장애인기업('26.06.30)"),('startup',"창업('26.06.30)")]
    for kind,sheet in mapping:
        for n,r in enumerate(w[sheet].iter_rows(min_row=2,values_only=True),2):
            if not any(r[:3]): continue
            add(kind,r[0],r[1],'policy',sheet,n,representative=r[2],start=r[5],end=r[6],cancel=r[7] if len(r)>7 else None)
        print(kind,len(registries[kind]),flush=True)
    sheet="중증('26.06.30)"
    for n,r in enumerate(w[sheet].iter_rows(min_row=2,values_only=True),2):
        if not any(r[:3]): continue
        add('severe',r[0],r[1],'policy',sheet,n,phone=r[6],address=r[5],start=r[8],end=r[9],facilityHead=clean(r[2]))
    sheet="장애인표준사업장('26.06.30"
    for n,r in enumerate(w[sheet].iter_rows(min_row=2,values_only=True),2):
        if not any(r[:3]): continue
        add('standard',r[0],r[1],'policy',sheet,n,representative=r[2],address=r[7],start=r[4],registeredWorkplaceName=clean(r[5]),registeredWorkplaceRepresentative=clean(r[6]),headOfficeBizno=bizno(r[3]),workplaceCertificationDate=date(r[8]))
    sheet="협동조합('26.06.30)"
    for n,r in enumerate(w[sheet].iter_rows(min_row=2,values_only=True),2):
        if not any(r[:2]): continue
        add('cooperative',r[0],r[1],'policy',sheet,n)
    w.close()
    w=openpyxl.load_workbook(social,read_only=True,data_only=True)
    sheet=w.active; it=sheet.iter_rows(values_only=True); next(it); headers=[clean(x) for x in next(it)]
    for n,rr in enumerate(it,3):
        r=dict(zip(headers,rr))
        add('social',r.get('사업자번호 (본점,지점)'),r.get('기업명'),'social',sheet.title,n,representative=r.get('대표자'),phone=r.get('대표전화번호'),address=r.get('소재지'),start=r.get('실적시작일(인증일자)'),end=r.get('실적종료일(실적인정)'),cancel=r.get('인증취소일(실적불인정)'),explicit=clean(r.get('인증상태')),certificationNo=clean(r.get('인증번호')),branch=clean(r.get('본지점')),sourceStatus=clean(r.get('인증상태')))
    w.close()
    summary={'generatedAt':dt.datetime.now(dt.timezone.utc).isoformat(),'evaluationDate':ASOF,'scope':'All source business records; no product or industry filters; historical records retained.','sources':{'policy':{'file':policy.name,'basisDate':'2026-06-30','sha256':hashlib.sha256(policy.read_bytes()).hexdigest()},'social':{'file':social.name,'basisDate':'2026-07-14','url':'https://www.sepp.or.kr/board/value/NOTICE/boardDtl?lettNo=278','sha256':hashlib.sha256(social.read_bytes()).hexdigest()}},'types':{}}
    def write(name,records):
        path=OUT/name; temp=path.with_suffix('.tmp')
        with temp.open('w',encoding='utf-8') as f:
            for row in records: f.write(json.dumps(row,ensure_ascii=False,separators=(',',':'))+'\n')
        temp.replace(path)
    for kind,companies in registries.items():
        counts=Counter()
        for code,c in companies.items():
            # Keep all source rows. Latest started record gives a summary, never live verification.
            begun=[r for r in c['records'] if not r['validFrom'] or r['validFrom']<=ASOF]
            latest=max(begun or c['records'],key=lambda r:(r['validFrom'] or '',r['cancelledAt'] if r['cancelledAt'] and r['cancelledAt']<=ASOF else '',r['validUntil'] or '',r['row']))
            c['name']=latest['name']; c['status']=latest['status']; counts[c['status']]+=1
            m=master.setdefault(code,{'bizno':code,'names':[],'types':[]})
            for r in c['records']:
                if r['name'] and r['name'] not in m['names']: m['names'].append(r['name'])
            m['types'].append({'type':TYPES[kind],'status':c['status'],'file':kind+'.jsonl'})
        write(kind+'.jsonl',(companies[k] for k in sorted(companies)))
        summary['types'][kind]={'label':TYPES[kind],'status':'extracted','sourceRows':rows[kind],'uniqueBusinessNumbers':len(companies),'invalidBusinessNumberRows':invalid[kind],'recordStatusCounts':dict(counts),'file':kind+'.jsonl'}
    summary['types']['sme']={'label':'중소기업','status':'account_permission_required','publicListTotalObserved':340219,'observedAt':ASOF,'url':'https://smpp.go.kr/cop/middlsslentrschtwr/selectMiddlSslentrSchtwrListVw.do','note':'Authenticated session and official menu verified: current account cannot access the service. Public legacy list lacks business numbers. No name-only joins.'}
    summary['uniqueBusinessNumbers']=len(master)
    write('businesses.jsonl',(master[k] for k in sorted(master)))
    (OUT/'manifest.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=True),flush=True)
    if previous_smpp:
        from import_smpp_sme import import_exports
        import_exports(ROOT/'private/sources/smpp'/previous_smpp['basisDate'],previous_smpp['basisDate'],previous['types']['sme']['publicListTotalObserved'])

if __name__=='__main__': run()
